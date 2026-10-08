"""One-off import of the hand-kept applications spreadsheet into the tracker (migration 004):
python -m jobpilot.import_applications [PATH] [--url-env NAME] [--dry-run]
PATH defaults to ~/Desktop/applications.xlsx; header row 1 holds Date, Company, Job title, Link, Score,
Status, Notes (any case, any order, extra columns ignored). Score is ignored: jobs has its own match_score.
Safe to re-run: (company, title, applied date) and the matched job are the duplicate keys (see 004).
All rows go in one transaction; --dry-run does the same work and rolls it back.
Role: run it with the owner/admin URL from .env (DATABASE_URL, or --url-env AZURE_DATABASE_URL). It's a
one-off run by the owner, not a service, so it needs no role of its own; jobpilot_feedback would also do."""
import argparse
import datetime as dt
import os
import re
import sys
from typing import Iterable, NamedTuple

from jobpilot import feedback
from jobpilot.db import connect, redact
from jobpilot.migrate import ROOT, add_url_env

try:
    from dotenv import load_dotenv
except ImportError:  # same fallback as job_searcher: env vars set by hand still work
    def load_dotenv(*args, **kwargs): return False

DEFAULT_PATH = os.path.join(os.path.expanduser("~"), "Desktop", "applications.xlsx")
COLUMNS = {"date": "date", "company": "company", "job title": "title", "link": "link",
           "status": "status", "notes": "notes"}  # header (lowercased) -> field
REQUIRED = ("company", "job title")
# words the spreadsheet used -> a 004 status; 'red' and 'ret' are Turkish for rejected
STATUS_WORDS = {
    "": "applied", "sent": "applied", "waiting": "applied", "pending": "applied", "başvuruldu": "applied",
    "test": "assessment", "coding test": "assessment", "online assessment": "assessment", "oa": "assessment",
    "case study": "assessment", "take-home": "assessment", "take home": "assessment",
    "interviewing": "interview", "interviewed": "interview", "mülakat": "interview",
    "rejected": "rejected", "rejection": "rejected", "red": "rejected", "ret": "rejected",
    "offered": "offer", "teklif": "offer", "withdrew": "withdrawn",
    **{s: s for s in feedback.STATUSES},
}
# /jobs/view/<id>, /jobs/view/<slug>-<id>, or ?currentJobId=<id> on search and collection pages
LINKEDIN_ID_RE = re.compile(r"linkedin\.com\S*?(?:/jobs/view/(?:[^/?#]*-)?(\d+)|[?&]currentJobId=(\d+))", re.I)


class Row(NamedTuple):
    line: int  # the Excel row number, so errors point at the right row
    applied_on: dt.date
    company: str
    title: str
    url: str | None
    status: str
    notes: str | None


def _text(value) -> str:
    if isinstance(value, float) and value.is_integer():
        value = int(value)  # a number typed into a text column: 2024.0 -> "2024"
    return "" if value is None else str(value).strip()


def parse_date(value, today: dt.date) -> dt.date:
    """Excel date, datetime, serial number or a 'YYYY-MM-DD' / 'DD.MM.YYYY' string; blank means today."""
    if isinstance(value, dt.datetime):  # check first: datetime is a subclass of date
        return value.date()
    if isinstance(value, dt.date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        from openpyxl.utils.datetime import from_excel  # a date cell without a date format stays a number
        try:
            return from_excel(value).date()
        except (ValueError, OverflowError, TypeError):
            raise ValueError(f"unreadable date {value!r}") from None
    text = _text(value)
    if not text:
        return today
    for fmt in ("%Y-%m-%d", "%d.%m.%Y"):
        try:
            return dt.datetime.strptime(text[:10], fmt).date()
        except ValueError:
            pass
    raise ValueError(f"unreadable date {text!r} (use YYYY-MM-DD or DD.MM.YYYY)")


def parse_status(value) -> str:
    word = " ".join(_text(value).lower().split())
    if word not in STATUS_WORDS:
        raise ValueError(f"unknown status {word!r} (use one of {', '.join(feedback.STATUSES)})")
    return STATUS_WORDS[word]


def linkedin_id(url: str | None) -> str | None:
    """The numeric LinkedIn job id in a link, which is how jobs.source_id stores LinkedIn jobs."""
    m = LINKEDIN_ID_RE.search(url or "")
    return (m.group(1) or m.group(2)) if m else None


def parse_rows(rows: Iterable[tuple], today: dt.date) -> tuple[list[Row], list[tuple[int, str]]]:
    """Header row + data rows (tuples of cell values) -> (rows, errors). Pure, so it's tested without a DB.
    A missing Company or Job title column raises ValueError; a bad row only becomes an (row number, reason)."""
    it = iter(rows)
    header = [" ".join(_text(h).lower().split()) for h in next(it, ())]  # "Job  Title " -> "job title"
    missing = [c for c in REQUIRED if c not in header]
    if missing:
        raise ValueError(f"missing column(s): {', '.join(missing)} (header row 1 needs Company and Job title)")
    pos = {COLUMNS[h]: i for i, h in reversed(list(enumerate(header))) if h in COLUMNS}  # first copy wins
    parsed, errors = [], []
    for line, cells in enumerate(it, start=2):
        cell = {f: (cells[i] if i < len(cells) else None) for f, i in pos.items()}
        if all(_text(c) == "" for c in cells):
            continue  # blank row, e.g. spacing in the sheet
        company, title = _text(cell["company"]), _text(cell["title"])
        try:
            if not company or not title:
                raise ValueError("Company and Job title are required")
            parsed.append(Row(line, parse_date(cell.get("date"), today), company, title,
                              _text(cell.get("link")) or None, parse_status(cell.get("status")),
                              _text(cell.get("notes")) or None))
        except ValueError as e:
            errors.append((line, str(e)))
    return parsed, errors


def read_sheet(path: str) -> list[tuple]:
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)  # data_only: formulas give their values
    try:
        return list(wb.worksheets[0].iter_rows(values_only=True))
    finally:
        wb.close()


def find_job(conn, url: str | None) -> int | None:
    """The JobPilot job behind a link: by LinkedIn id first (survives tracking parameters), else exact url."""
    if not url:
        return None
    sid = linkedin_id(url)
    if sid:
        row = conn.execute("SELECT id FROM jobs WHERE source = 'linkedin' AND source_id = %s", (sid,)).fetchone()
        if row:
            return row[0]
    row = conn.execute("SELECT min(id) FROM jobs WHERE apply_url = %s", (url,)).fetchone()
    return row[0]


def import_rows(conn, rows: list[Row], dry_run: bool = False, log=print) -> dict[str, int]:
    """Writes the rows in one transaction (rolled back for a dry run); returns the counts."""
    counts = {"inserted": 0, "present": 0, "matched": 0}
    with conn.transaction(force_rollback=dry_run):
        for r in rows:
            job_id = find_job(conn, r.url)
            app_id, created = feedback.import_application(conn, r.company, r.title, r.url, r.applied_on,
                                                          r.notes, job_id)
            # only on a new row: a re-run must not overwrite a status the bot has moved on since
            if created and r.status != "applied":
                feedback.set_status(conn, app_id, r.status)
            counts["inserted" if created else "present"] += 1
            counts["matched"] += job_id is not None
            if dry_run:
                what = "insert" if created else "skip (already present)"
                job = f", job {job_id}" if job_id is not None else ""
                log(f"  row {r.line}: would {what}: {r.company} / {r.title}, {r.applied_on}, {r.status}{job}")
    return counts


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(prog="python -m jobpilot.import_applications",
                                     description="Import the applications spreadsheet into the tracker")
    parser.add_argument("path", nargs="?", default=DEFAULT_PATH, help=f"the .xlsx file (default: {DEFAULT_PATH})")
    parser.add_argument("--dry-run", action="store_true", help="show what would be imported, then roll back")
    add_url_env(parser)
    args = parser.parse_intermixed_args(sys.argv[1:] if argv is None else argv)
    if not os.path.exists(args.path):
        print(f"ERROR: no such file: {args.path}", file=sys.stderr)
        return 1
    try:
        rows, errors = parse_rows(read_sheet(args.path), dt.date.today())
    except Exception as e:  # a missing column, or not an xlsx file
        print(f"ERROR: can't read {os.path.basename(args.path)}: {type(e).__name__}: {e}", file=sys.stderr)
        return 1
    load_dotenv(os.path.join(ROOT, ".env"))
    url = os.environ.get(args.url_env, "").strip()
    if not url:
        print(f"ERROR: {args.url_env} not set", file=sys.stderr)  # the name only, never a value
        return 1
    try:
        with connect(url) as conn:
            counts = import_rows(conn, rows, args.dry_run)
    except Exception as e:  # nothing was written: the one transaction rolled back
        print(f"ERROR: import failed, nothing written: {redact(e, url)}", file=sys.stderr)
        return 1
    print(f"{'DRY RUN (rolled back): ' if args.dry_run else ''}{len(rows) + len(errors)} rows read, "
          f"{counts['inserted']} inserted, {counts['present']} already present, "
          f"{counts['matched']} matched to a job, {len(errors)} errors")
    for line, reason in errors:
        print(f"  row {line}: {reason}")
    if errors:  # the good rows are in; fixing these and re-running won't duplicate them
        print("Fix the rows above and re-run; rows already imported are skipped.")
    return 1 if errors else 0


if __name__ == "__main__":
    sys.exit(main())
